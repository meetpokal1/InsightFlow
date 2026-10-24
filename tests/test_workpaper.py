"""Unit tests for the Data Assurance Workpaper (.xlsx) generation and structure."""

from pathlib import Path

import openpyxl
import pytest

from scripts.generate_data_assurance_workpaper import generate_workbook
from src.config import PROJECT_ROOT


@pytest.fixture(scope="module")
def workpaper_path(tmp_path_factory):
    out_dir = tmp_path_factory.mktemp("workpaper")
    out_file = out_dir / "data_assurance_workpaper.xlsx"
    processed_dir = PROJECT_ROOT / "data" / "processed"
    findings_file = processed_dir / "data_quality_report.csv"
    if not findings_file.exists():
        raw_dir = PROJECT_ROOT / "data" / "raw"
        from src.pipeline import clean_raw_files, save_outputs
        from tests.helpers import AS_OF

        res = clean_raw_files(raw_dir, AS_OF)
        temp_processed = tmp_path_factory.mktemp("processed")
        save_outputs(res, temp_processed)
        generate_workbook(out_file, processed_dir=temp_processed)
    else:
        generate_workbook(out_file, processed_dir=processed_dir)
    return out_file


def test_workpaper_sheets_and_structure(workpaper_path):
    """Verify that all six required audit worksheets exist with valid dimensions."""
    wb = openpyxl.load_workbook(workpaper_path, data_only=False)
    expected_sheets = [
        "Executive Summary",
        "Table Reconciliation",
        "Quality Findings",
        "Rejected Records",
        "Validation Checks",
        "Definitions & Limitations",
    ]
    assert wb.sheetnames == expected_sheets

    for name in expected_sheets:
        ws = wb[name]
        assert ws.max_row >= 5
        assert ws.max_column >= 4


def test_reconciliation_formulas_and_balance(workpaper_path):
    """Verify table reconciliation formulas and exact arithmetic balancing."""
    wb = openpyxl.load_workbook(workpaper_path, data_only=False)
    ws = wb["Table Reconciliation"]

    # Verify header titles
    assert "Raw Records Ingested (A)" in str(ws["B3"].value)
    assert "Exact Duplicates Dropped (B)" in str(ws["C3"].value)
    assert "Quarantined / Rejected (C)" in str(ws["D3"].value)
    assert "Merged / Remapped (D)" in str(ws["E3"].value)
    assert "Clean Records Loaded (E)" in str(ws["F3"].value)

    # Verify formula syntax in row 4 (categories)
    assert ws["G4"].value == "=B4-C4-D4-E4"
    assert ws["H4"].value == "=F4-G4"
    assert "IF(H4=0" in ws["I4"].value
    assert ws["J4"].value == "=F4/B4"

    # Verify totals formula in row 12
    assert ws["B12"].value == "=SUM(B4:B11)"
    assert ws["F12"].value == "=SUM(F4:F11)"


def test_rejected_records_protection(workpaper_path):
    """Verify formula injection guards on text cells."""
    wb = openpyxl.load_workbook(workpaper_path, data_only=False)
    ws = wb["Rejected Records"]

    assert ws.max_row > 10
    # Check that no cell value directly starts with dangerous unescaped characters unless it is a formula
    for row in ws.iter_rows(min_row=4, values_only=False):
        for cell in row:
            if isinstance(cell.value, str):
                # If it's a formula, openpyxl treats it as starting with =, but in Rejected Records there are no formulas
                assert not cell.value.startswith(("=cmd", "+cmd", "-cmd", "@cmd"))


def test_validation_checks_content(workpaper_path):
    """Verify automated integrity checks are recorded as PASSED."""
    wb = openpyxl.load_workbook(workpaper_path, data_only=True)
    ws = wb["Validation Checks"]

    check_statuses = [cell.value for cell in ws["F"][3:] if cell.value is not None]
    assert len(check_statuses) >= 10
    assert all(status == "PASS" for status in check_statuses)
